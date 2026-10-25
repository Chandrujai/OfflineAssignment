import os
import io
import re
import streamlit as st
import networkx as nx
import matplotlib.pyplot as plt
import pypdf
from docx import Document as DocxDocument
from dotenv import load_dotenv

# Load environment variables from .env
load_dotenv()

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
NEO4J_URI = os.getenv("NEO4J_URI", "bolt://localhost:7687")
NEO4J_USERNAME = os.getenv("NEO4J_USERNAME") or os.getenv("NEO4J_USER", "neo4j")
NEO4J_PASSWORD = os.getenv("NEO4J_PASSWORD", "password")
NEO4J_DATABASE = os.getenv("NEO4J_DATABASE", "neo4j")
AURA_INSTANCEID = os.getenv("AURA_INSTANCEID", "")
AURA_INSTANCENAME = os.getenv("AURA_INSTANCENAME", "")

if OPENAI_API_KEY:
    os.environ["OPENAI_API_KEY"] = OPENAI_API_KEY

# LangChain & Modern Package Imports (Resolved Deprecations)
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings, ChatOpenAI
from langchain_chroma import Chroma
from langchain_neo4j import Neo4jGraph, LLMGraphTransformer
from neo4j import GraphDatabase
from neo4j.exceptions import ServiceUnavailable, AuthError

# ==============================================================================
# Page Configuration
# ==============================================================================
st.set_page_config(
    page_title="Hybrid RAG: Vector DB + Neo4j Graph RAG",
    page_icon="🕸️",
    layout="wide"
)

# ==============================================================================
# Session State Initialization
# ==============================================================================
if "messages" not in st.session_state:
    st.session_state.messages = []
if "vector_db" not in st.session_state:
    st.session_state.vector_db = None
if "indexed" not in st.session_state:
    st.session_state.indexed = False

# ==============================================================================
# Helper Functions: Robust File Parsers
# ==============================================================================
def parse_uploaded_file(uploaded_file):
    temp_filename = f"temp_{uploaded_file.name}"
    with open(temp_filename, "wb") as f:
        f.write(uploaded_file.getbuffer())
    
    docs = []
    file_ext = uploaded_file.name.split(".")[-1].lower()
    
    try:
        if file_ext == "pdf":
            reader = pypdf.PdfReader(temp_filename)
            text = ""
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text += f"\n--- Page {i+1} ---\n" + page_text
            if text.strip():
                docs.append(Document(page_content=text, metadata={"source": uploaded_file.name}))
                
        elif file_ext in ["docx", "doc"]:
            doc = DocxDocument(temp_filename)
            text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            if text.strip():
                docs.append(Document(page_content=text, metadata={"source": uploaded_file.name}))
                
        elif file_ext in ["txt", "md"]:
            with open(temp_filename, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
            if text.strip():
                docs.append(Document(page_content=text, metadata={"source": uploaded_file.name}))
    finally:
        if os.path.exists(temp_filename):
            os.remove(temp_filename)
            
    return docs

# ==============================================================================
# Helper Functions: Graph Queries
# ==============================================================================
def query_neo4j_graph(neo4j_graph, query_text):
    keywords = re.findall(r'\b[A-Z][a-zA-Z0-9_]{2,}\b', query_text)
    if not keywords:
        keywords = [w for w in query_text.split() if len(w) > 3]

    cypher_query = """
    MATCH (n)-[r]->(m)
    WHERE any(kw IN $keywords WHERE toLower(n.id) CONTAINS toLower(kw) OR toLower(m.id) CONTAINS toLower(kw))
    RETURN n.id AS source, type(r) AS relation, m.id AS target
    LIMIT 25
    """
    try:
        results = neo4j_graph.query(cypher_query, params={"keywords": keywords})
        return results
    except Exception as e:
        st.warning(f"Graph query execution warning: {e}")
        return []

def fetch_graph_networkx(neo4j_graph):
    cypher_query = """
    MATCH (n)-[r]->(m)
    RETURN n.id AS source, labels(n)[0] AS source_label, type(r) AS relation, m.id AS target, labels(m)[0] AS target_label
    LIMIT 100
    """
    try:
        results = neo4j_graph.query(cypher_query)
        G = nx.DiGraph()
        for row in results:
            src = row["source"]
            tgt = row["target"]
            rel = row["relation"]
            G.add_node(src, label=row.get("source_label", "Entity"))
            G.add_node(tgt, label=row.get("target_label", "Entity"))
            G.add_edge(src, tgt, label=rel)
        return G
    except Exception as e:
        st.error(f"Failed to fetch graph for visualization: {e}")
        return nx.DiGraph()

# ==============================================================================
# UI Layout: Sidebar Configuration
# ==============================================================================
st.title("🧠 Hybrid RAG Studio (Local Testing)")

with st.sidebar:
    st.header("System Status")
    st.write(f"- **OpenAI API Key Configured:** {'✅ Yes' if OPENAI_API_KEY else '❌ No'}")
    st.write(f"- **Neo4j Target URI:** `{NEO4J_URI}`")
    st.write(f"- **Neo4j Database:** `{NEO4J_DATABASE}`")
    if AURA_INSTANCENAME:
        st.write(f"- **Aura Instance:** `{AURA_INSTANCENAME}`")

    st.divider()
    st.header("Document Ingestion")
    uploaded_files = st.file_uploader(
        "Upload Documents (.pdf, .docx, .txt, .md)",
        type=["pdf", "docx", "doc", "txt", "md"],
        accept_multiple_files=True
    )
    
    chunk_size = st.slider("Chunk Size", min_value=200, max_value=2000, value=800, step=100)
    chunk_overlap = st.slider("Chunk Overlap", min_value=0, max_value=400, value=100, step=20)
    
    process_button = st.button("🚀 Process & Index Documents")

# ==============================================================================
# Document Ingestion Pipeline Orchestration
# ==============================================================================
if process_button:
    if not OPENAI_API_KEY:
        st.error("OpenAI API Key is missing. Check your .env configuration file.")
    elif not uploaded_files:
        st.warning("Please upload at least one document.")
    else:
        with st.spinner("Parsing documents & building Hybrid Vector/Graph Indices..."):
            try:
                all_docs = []
                for file in uploaded_files:
                    docs = parse_uploaded_file(file)
                    all_docs.extend(docs)
                
                if not all_docs:
                    st.error("No extractable text found in the uploaded files.")
                else:
                    text_splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
                    chunks = text_splitter.split_documents(all_docs)
                    
                    # Vector DB
                    embeddings = OpenAIEmbeddings(model="text-embedding-3-small")
                    vector_db = Chroma.from_documents(
                        documents=chunks,
                        embedding=embeddings,
                        persist_directory="./chroma_db_store"
                    )
                    st.session_state.vector_db = vector_db
                    
                    # Neo4j Graph (Using updated environment variables & database)
                    neo4j_graph = Neo4jGraph(
                        url=NEO4J_URI,
                        username=NEO4J_USERNAME,
                        password=NEO4J_PASSWORD,
                        database=NEO4J_DATABASE
                    )
                    
                    llm = ChatOpenAI(model="gpt-4o-mini", temperature=0)
                    transformer = LLMGraphTransformer(llm=llm)
                    graph_docs = transformer.convert_to_graph_documents(chunks)
                    
                    # Fixed parameter: baseEntityLabel (camelCase)
                    neo4j_graph.add_graph_documents(graph_docs, baseEntityLabel=True, include_source=True)
                    
                    st.session_state.indexed = True
                    st.session_state.neo4j_graph = neo4j_graph
                    st.success(f"Processing Complete! Created {len(chunks)} vector chunks & updated Neo4j Knowledge Graph.")
                
            except Exception as e:
                import traceback
                st.error(f"Error during ingestion pipeline: {e}")
                st.code(traceback.format_exc())
                print(traceback.format_exc(), flush=True)

# ==============================================================================
# UI Navigation Tabs
# ==============================================================================
tab_chat, tab_graph, tab_info = st.tabs(["💬 Hybrid RAG Chat", "🕸️ Interactive Graph Explorer", "ℹ️ System Info"])

with tab_chat:
    st.markdown("### Ask Questions Across Vector Context and Knowledge Graph")
    
    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    user_query = st.chat_input("Ask a question about your documents...")
    if user_query:
        st.session_state.messages.append({"role": "user", "content": user_query})
        with st.chat_message("user"):
            st.markdown(user_query)

        if not st.session_state.get("indexed", False):
            response = "⚠️ Please upload and index documents using the sidebar first."
            st.session_state.messages.append({"role": "assistant", "content": response})
            with st.chat_message("assistant"):
                st.markdown(response)
        else:
            with st.chat_message("assistant"):
                with st.spinner("Retrieving semantic & structural knowledge..."):
                    vector_results = st.session_state.vector_db.similarity_search(user_query, k=3)
                    vector_context = "\n\n".join([f"Chunk {i+1}:\n" + doc.page_content for i, doc in enumerate(vector_results)])

                    neo4j_results = query_neo4j_graph(st.session_state.neo4j_graph, user_query)
                    graph_context = ""
                    if neo4j_results:
                        graph_tuples = [f"({row['source']}) -[{row['relation']}]-> ({row['target']})" for row in neo4j_results]
                        graph_context = "\n".join(graph_tuples)
                    else:
                        graph_context = "No direct entity-relationship triplets matched in the Knowledge Graph."

                    hybrid_prompt = f"""
You are an advanced AI assistant utilizing a Hybrid RAG framework. Synthesize a comprehensive answer using BOTH the vector similarity chunks and structural Knowledge Graph triplets provided below.

### 1. Vector Semantic Context (Document Chunks):
{vector_context}

### 2. Knowledge Graph Context (Extracted Entity Relationships):
{graph_context}

### User Query:
{user_query}

### Formulated Answer:
"""
                    llm = ChatOpenAI(model="gpt-4o-mini", temperature=0.2)
                    response_obj = llm.invoke(hybrid_prompt)
                    final_answer = response_obj.content

                    st.markdown(final_answer)

                    with st.expander("🔍 View Retrieved Hybrid Context Details"):
                        st.subheader("Semantic Vector Chunks")
                        st.text(vector_context)
                        st.subheader("Knowledge Graph Triplets")
                        st.text(graph_context)

                    st.session_state.messages.append({"role": "assistant", "content": final_answer})

with tab_graph:
    st.markdown("### Neo4j Knowledge Graph Visualizer")
    if st.session_state.get("indexed", False) and "neo4j_graph" in st.session_state:
        if st.button("🔄 Refresh Graph View"):
            st.rerun()
            
        G = fetch_graph_networkx(st.session_state.neo4j_graph)
        if len(G.nodes) > 0:
            fig, ax = plt.subplots(figsize=(12, 7))
            pos = nx.spring_layout(G, k=0.3, seed=42)
            nx.draw_networkx_nodes(G, pos, node_size=700, node_color="#4CAF50", alpha=0.9, ax=ax)
            nx.draw_networkx_edges(G, pos, width=1.5, alpha=0.5, edge_color="gray", ax=ax)
            nx.draw_networkx_labels(G, pos, font_size=8, font_weight="bold", ax=ax)
            edge_labels = nx.get_edge_attributes(G, 'label')
            nx.draw_networkx_edge_labels(G, pos, edge_labels=edge_labels, font_size=6, ax=ax)
            plt.title("Extracted Neo4j Entities and Relationships", fontsize=12)
            plt.axis("off")
            st.pyplot(fig)
        else:
            st.info("The Knowledge Graph currently contains no nodes.")
    else:
        st.info("Please complete the document ingestion step in the sidebar to render the graph visualization.")

with tab_info:
    st.markdown("### System Configuration Status")
    st.write(f"- **OpenAI API Key Configured:** {'✅ Yes' if OPENAI_API_KEY else '❌ No'}")
    st.write(f"- **Neo4j URI:** `{NEO4J_URI}`")
    st.write(f"- **Neo4j Username:** `{NEO4J_USERNAME}`")
    st.write(f"- **Neo4j Database:** `{NEO4J_DATABASE}`")
    st.write(f"- **Aura Instance ID:** `{AURA_INSTANCEID}`")
    st.write(f"- **Aura Instance Name:** `{AURA_INSTANCENAME}`")
    st.write(f"- **Documents Indexed:** {'✅ Yes' if st.session_state.get('indexed') else '❌ No'}")